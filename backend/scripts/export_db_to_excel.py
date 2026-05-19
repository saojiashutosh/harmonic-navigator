"""Export the entire Harmonic Navigator DB to a single Excel workbook.

One sheet per model (named `<app>.<Model>`), header row uses the Django
`attname` for each field (so foreign keys show up as `<field>_id`). UUID PKs
and FK references are written as strings; datetimes are written as ISO
strings; JSON fields are JSON-encoded.

Designed to be paired with `import_db_from_excel.py` so a new contributor
can clone the repo, bring up `docker-compose`, and load the same data the
maintainer has.

Usage (inside the `web` container — recommended, since Django + Postgres
are already wired up there):

    docker compose exec web python scripts/export_db_to_excel.py
    docker compose exec web python scripts/export_db_to_excel.py -o /app/data/full_db_export.xlsx

Or directly on the host if your env is set up:

    python backend/scripts/export_db_to_excel.py -o backend/data/full_db_export.xlsx
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from uuid import UUID

BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "core.settings.development")

import django  # noqa: E402

django.setup()

from django.apps import apps  # noqa: E402
from django.db import models as dj_models  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402


# Order matters for the matching import script: parents before children so
# the importer can satisfy FK references row-by-row.
MODEL_ORDER: list[tuple[str, str]] = [
    ("users", "Users"),
    ("users", "UsersDevices"),
    ("tracks", "Artist"),
    ("tracks", "MoodTag"),
    ("tracks", "Track"),
    ("tracks", "TrackMoodTag"),
    ("tracks", "AudioFeatureSnapshot"),
    ("moods", "Question"),
    ("moods", "MoodSession"),
    ("moods", "Answer"),
    ("moods", "MoodInference"),
    ("playlists", "Playlist"),
    ("playlists", "PlaylistTrack"),
    ("playlists", "SavedPlaylist"),
    ("groups", "GroupSession"),
    ("groups", "GroupParticipant"),
    ("feedback", "TrackFeedback"),
    ("feedback", "UserMoodPreference"),
    ("feedback", "TrackMoodScore"),
    ("music", "MusicTrack"),
]


HEADER_FILL = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
HEADER_FONT = Font(color="E0C97F", bold=True)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


def _sheet_name(app_label: str, model_name: str) -> str:
    # Excel caps sheet names at 31 chars.
    return f"{app_label}.{model_name}"[:31]


def _serialize(field, value):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(field, dj_models.JSONField):
        return json.dumps(value, default=str, ensure_ascii=False)
    return value


def _export_model(wb: Workbook, app_label: str, model_name: str) -> int:
    Model = apps.get_model(app_label, model_name)
    ws = wb.create_sheet(_sheet_name(app_label, model_name))

    fields = list(Model._meta.fields)
    headers = [f.attname for f in fields]  # FKs become `<name>_id`
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN

    count = 0
    for obj in Model.objects.all().order_by("pk").iterator():
        row = []
        for f in fields:
            raw = getattr(obj, f.attname)
            row.append(_serialize(f, raw))
        ws.append(row)
        count += 1

    # Light auto-sizing
    for col in ws.columns:
        max_len = max((len(str(cell.value)) if cell.value is not None else 0 for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 60)

    return count


def main() -> None:
    parser = argparse.ArgumentParser(description="Export the entire DB to an Excel workbook.")
    parser.add_argument(
        "-o", "--output",
        default="/app/data/full_db_export.xlsx",
        help="Destination Excel path (default: /app/data/full_db_export.xlsx)",
    )
    args = parser.parse_args()

    wb = Workbook()
    # Drop the default empty sheet
    wb.remove(wb.active)

    total = 0
    print("Exporting models:")
    for app_label, model_name in MODEL_ORDER:
        try:
            apps.get_model(app_label, model_name)
        except LookupError:
            print(f"  - {app_label}.{model_name}: SKIPPED (model not found)")
            continue
        n = _export_model(wb, app_label, model_name)
        total += n
        print(f"  - {app_label}.{model_name}: {n} rows")

    out = args.output
    out_dir = os.path.dirname(out)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(out)
    print(f"\nWrote {total} rows across {len(wb.sheetnames)} sheets to {out}")


if __name__ == "__main__":
    main()
