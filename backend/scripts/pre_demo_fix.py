"""One-shot script that re-applies on the restored xlsx all the data fixes
that had accumulated across the demo-prep session, but were lost when an
incompatible export_tracks_to_excel call rewrote the file in a different
format.

Operations:

1. NULL release_year for tracks where it's clearly a Spotify
   compilation/remaster year (title contains Jhankar/JB/Lofi/etc.) or
   the artist is unambiguously a pre-2005 golden-era composer/singer
   (Kalyanji-Anandji, Madan Mohan, Naushad, etc.) and the year is >=2010.

2. Re-derive primary_mood with the smarter classifier (title-keyword
   overrides -> genre floors -> looser feature thresholds), matching
   the retag_moods management command.

Run from the project root:
    python backend/scripts/pre_demo_fix.py
"""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

XLSX_PATH = Path(__file__).resolve().parents[1] / "data" / "harmonic_export.xlsx"

# ─────────────────────────────────────────────────────────────────────────────
# 1. ERA / RELEASE YEAR FIX
# ─────────────────────────────────────────────────────────────────────────────

REMIX_PAT = re.compile(
    r"\bJhankar\b|\bJB\b|\bLofi\b|\bSlowed\b|\bReverb\b|\bRemaster|"
    r"\bReprise\b|\bUnplugged\b|\bRevisited\b|\bRecreated\b|Lo[\s\-]?Fi",
    re.IGNORECASE,
)

# Pre-2005-peak artists. Any post-2010 release_year on these is almost
# certainly a Spotify-side compilation or re-release tag, not the original.
CLASSIC_ARTISTS = {
    # Original list (from previous era-fix commit)
    "Kumar Sanu",
    "K. S. Chithra", "Chitra",
    "Udit Narayan",
    "Alka Yagnik",
    "Jatin-Lalit",
    "Anand-Milind",
    "Anu Malik",
    "Bappi Lahiri",
    "R.D. Burman", "RD Burman",
    "Asha Bhosle",
    "Lata Mangeshkar",
    "Mohammed Rafi", "Mohd. Rafi",
    "Kishore Kumar",
    "Mukesh",
    "Pankaj Udhas",
    "Vinod Rathod",
    "Kavita Krishnamurthy",
    "Hariharan",
    "Suresh Wadkar",
    "Lucky Ali",
    "Sayeed Quadri",
    "Aadesh Shrivastava",
    "Roop Kumar Rathod",
    "Labh Janjua",
    "Talat Mahmood",
    "Shamshad Begum",
    "Manna Dey",
    "Hemant Kumar",
    "Geeta Dutt",
    "Suman Kalyanpur",
    "Mahendra Kapoor",
    "Anuradha Paudwal",
    "Sadhana Sargam",
    "Poornima",
    "Nitin Mukesh",
    "Abhijeet", "Abhijeet Bhattacharya",
    "Babul Supriyo",
    "Sapna Mukherjee",
    "Sukhwinder Singh",
    # Newly added composers/singers from golden-era audit
    "Kalyanji - Anandji", "Kalyanji-Anandji",
    "Anand Raaj Anand", "Anand Raj Anand",
    "Madan Mohan",
    "C. Ramchandra",
    "Naushad",
    "Roshan", "Rajesh Roshan",
    "S.D. Burman", "SD Burman",
    "Salil Choudhury", "Salil Chowdhury",
    "O.P. Nayyar", "OP Nayyar",
    "Laxmikant-Pyarelal",
    "Shankar Jaikishan",
    "Ravindra Jain",
    "Chitragupta",
    "Hridaynath Mangeshkar",
    "Khayyam",
    "Nusrat Fateh Ali Khan",
    "Mehdi Hassan",
    "Jagjit Singh",
    "Chitra Singh",
    "Anand Shinde",
}


def apply_era_fix(ws) -> tuple[int, int]:
    """NULL out remaster/compilation years and post-2010 classic-artist years."""
    header = [c.value for c in ws[1]]
    ti, ai, yi = header.index("title"), header.index("artist_name"), header.index("release_year")
    n_remix = n_classic = 0

    for row in ws.iter_rows(min_row=2):
        title, artist, year = row[ti].value, row[ai].value, row[yi].value
        if not title or not artist or year is None or year < 2010:
            continue
        if REMIX_PAT.search(title):
            row[yi].value = None
            n_remix += 1
        elif artist in CLASSIC_ARTISTS:
            row[yi].value = None
            n_classic += 1

    return n_remix, n_classic


# ─────────────────────────────────────────────────────────────────────────────
# 2. MOOD RETAG
# ─────────────────────────────────────────────────────────────────────────────

MELANCHOLIC_WORDS = re.compile(
    r"\b(sad|tanha|tanhai|bewafa|judaai|dard|udaas|alvida|bichd|adhoor|"
    r"akela|gum|gham|akele|toot|bhula|yaad|rona|broken|lonely|tears|"
    r"misery|cry|alone|farewell|goodbye)\b",
    re.IGNORECASE,
)
CELEBRATORY_WORDS = re.compile(
    r"\b(naach|dance|party|jhoom|jhoome|dhoom|baraat|baaraat|sangeet|"
    r"masti|dhamaal|dhamaka|zingaat|wedding|celebrat|disco|jashn|raas|"
    r"rangrez|holi|garba|dandiya|bhangra|jingle)\b",
    re.IGNORECASE,
)
CALM_WORDS = re.compile(
    r"\b(lullaby|sleep|nidra|raat|chaand|sukoon|shanti|peace|meditation|"
    r"prayer|bhajan|aarti|mantra|stotra|stuti|kirtan|saregama)\b",
    re.IGNORECASE,
)

CELEBRATORY_GENRES = {"bhangra", "edm", "dance", "disco", "funk"}
CALM_GENRES = {"devotional", "spiritual", "bhajan", "lullaby", "ambient", "lo-fi", "lofi"}


def _g(value) -> str:
    return (value or "").strip().lower() if isinstance(value, str) else ""


def derive_mood(title, genre, energy, valence, tempo, language, is_instrumental) -> str:
    title = title or ""
    genre = _g(genre)
    if MELANCHOLIC_WORDS.search(title):
        return "melancholic"
    if CELEBRATORY_WORDS.search(title):
        return "celebratory"
    if CALM_WORDS.search(title):
        return "calm"
    if genre in CALM_GENRES:
        return "calm"
    if genre in CELEBRATORY_GENRES and tempo and tempo >= 115:
        return "celebratory"
    if energy is not None and valence is not None:
        if energy >= 0.70 and valence >= 0.60:
            return "celebratory"
        if energy >= 0.60 and valence >= 0.45:
            return "energized"
        if energy <= 0.40 and valence <= 0.40:
            return "melancholic"
        if energy <= 0.55 and valence >= 0.30:
            return "calm"
        return "focused"
    lang = _g(language)
    if tempo:
        if tempo >= 125 and genre in {"bollywood", "marathi", "punjabi", "pop", "rock"}:
            return "celebratory"
        if tempo >= 110:
            return "energized"
        if tempo <= 75:
            return "calm"
    if genre in {"bollywood", "marathi", "pop", "rock"}:
        return "energized"
    if lang == "instrumental" or is_instrumental:
        return "focused"
    return "calm"


def apply_mood_retag(ws) -> tuple[int, Counter, Counter]:
    header = [c.value for c in ws[1]]
    idx = {name: header.index(name) for name in (
        "title", "genre", "energy", "valence", "tempo_bpm", "language",
        "is_instrumental", "primary_mood",
    )}
    before = Counter()
    after = Counter()
    changed = 0
    for row in ws.iter_rows(min_row=2):
        if all(c.value is None for c in row):
            continue
        cur = row[idx["primary_mood"]].value
        before[cur] += 1
        new = derive_mood(
            title=row[idx["title"]].value,
            genre=row[idx["genre"]].value,
            energy=row[idx["energy"]].value,
            valence=row[idx["valence"]].value,
            tempo=row[idx["tempo_bpm"]].value,
            language=row[idx["language"]].value,
            is_instrumental=row[idx["is_instrumental"]].value,
        )
        after[new] += 1
        if cur != new:
            row[idx["primary_mood"]].value = new
            changed += 1
    return changed, before, after


# ─────────────────────────────────────────────────────────────────────────────

def main():
    print(f"Opening {XLSX_PATH}")
    wb = load_workbook(XLSX_PATH)
    ws = wb["tracks"]

    n_remix, n_classic = apply_era_fix(ws)
    print(f"\nEra fix:")
    print(f"  remix/remaster-marker tracks nulled: {n_remix}")
    print(f"  classic-artist post-2010 tracks nulled: {n_classic}")

    changed, before, after = apply_mood_retag(ws)
    print(f"\nMood retag — {changed} tracks updated:")
    print(f"  before: {dict(before.most_common())}")
    print(f"  after:  {dict(after.most_common())}")

    wb.save(XLSX_PATH)
    print(f"\nSaved {XLSX_PATH}")


if __name__ == "__main__":
    main()
