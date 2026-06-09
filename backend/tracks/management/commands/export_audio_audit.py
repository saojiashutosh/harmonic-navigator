"""Export an audit-quality Excel workbook of every analysed Track.

Three sheets:

  * **Tracks** — every active Track with the full feature set, one row each.
  * **Snapshots** — the latest :class:`AudioFeatureSnapshot` per track,
    flattened so the underlying DSP feature vector is visible alongside the
    Track-level fields.
  * **Summary** — distribution stats (mean/median/stdev/range) per feature
    plus counts by source, language, mood, genre.

This is the canonical artefact for reviewing what the engine produced. The
workbook lands in ``backend/data/exports/audio_audit_<timestamp>.xlsx``.

Usage (inside Docker):

    python manage.py export_audio_audit
    python manage.py export_audio_audit --path /tmp/my_audit.xlsx
    python manage.py export_audio_audit --include-snapshots false
"""
from __future__ import annotations

import statistics as stats
from collections import Counter
from datetime import datetime
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand

from tracks.models import AudioFeatureSnapshot, Track


_NUMERIC_FEATURES = (
    "energy", "valence", "acousticness", "instrumentalness",
    "loudness", "tempoBpm",
)

# Snapshot fields worth surfacing — JSON column lookups would be ugly, so
# we flatten them on export.
_SNAPSHOT_DSP_KEYS = (
    "rms_mean", "rms_std", "dynamic_range", "loudness_db",
    "brightness", "spectral_rolloff", "spectral_bandwidth",
    "spectral_flatness", "spectral_contrast", "zero_crossing_rate",
    "tempo_bpm", "beat_strength", "onset_rate", "pulse_clarity",
    "harmonic_ratio", "percussive_ratio", "low_freq_ratio",
    "vocal_band_ratio", "key", "mode", "key_strength",
    "duration_sec", "sample_rate",
)


class Command(BaseCommand):
    help = "Export an audit-quality Excel of every analysed track."

    def add_arguments(self, parser):
        parser.add_argument("--path", default=None)
        parser.add_argument(
            "--include-snapshots",
            type=lambda v: str(v).lower() in {"1", "true", "yes"},
            default=True,
            help="Include the Snapshots sheet (slower if you have many snapshots).",
        )
        parser.add_argument(
            "--only-analysed",
            action="store_true",
            help="Skip tracks with no analysis yet (default: include).",
        )

    # ------------------------------------------------------------------
    def handle(self, *args, **options):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        only_analysed = options["only_analysed"]
        include_snapshots = options["include_snapshots"]

        qs = Track.objects.filter(isActive=True).select_related("artistId")
        if only_analysed:
            qs = qs.filter(
                energy__isnull=False,
                valence__isnull=False,
                tempoBpm__isnull=False,
            )

        tracks = list(qs.order_by("language", "source", "title"))
        self.stdout.write(self.style.MIGRATE_HEADING(
            f"Exporting {len(tracks)} tracks "
            f"(snapshots={'YES' if include_snapshots else 'NO'})"
        ))

        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2C3E50")
        center = Alignment(horizontal="center")

        # -------------------- Tracks sheet --------------------
        ws = wb.active
        ws.title = "Tracks"
        ws.freeze_panes = "C2"  # freeze id+title cols + header row
        track_headers = [
            "id", "title", "artist", "language", "source", "release_year",
            "duration_ms",
            "energy", "valence", "tempo_bpm", "acousticness",
            "instrumentalness", "loudness", "key_signature",
            "primary_mood", "genre", "region", "is_instrumental",
            "is_explicit", "artist_popularity",
            "stream_url", "spotify_id",
            "features_synced_at", "created_at", "updated_at",
        ]
        ws.append(track_headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        for t in tracks:
            ws.append([
                str(t.id), t.title or "",
                (t.artistId.name if t.artistId else "") or "",
                t.language or "", t.source or "", t.releaseYear,
                t.durationMs,
                t.energy, t.valence, t.tempoBpm, t.acousticness,
                t.instrumentalness, t.loudness, t.keySignature,
                t.primaryMood, t.genre, t.region,
                t.isInstrumental, t.isExplicit, t.artistPopularity,
                (t.streamUrl or "")[:200],
                t.spotifyId or "",
                _fmt_dt(t.featuresSyncedAt),
                _fmt_dt(t.createdAt),
                _fmt_dt(t.updatedAt),
            ])
        self._autofit(ws, track_headers, get_column_letter)

        # -------------------- Snapshots sheet --------------------
        if include_snapshots:
            ws2 = wb.create_sheet("Snapshots")
            ws2.freeze_panes = "B2"
            snap_headers = [
                "track_id", "title", "snapshot_id", "synced_at",
                "calibrated",
                # top-level engine outputs
                "energy", "valence", "tempo_bpm", "acousticness",
                "instrumentalness", "loudness", "key_signature",
                "primary_mood", "mood_confidence",
                "genre", "genre_confidence",
                "region", "region_confidence",
                "duration_sec",
            ] + list(_SNAPSHOT_DSP_KEYS) + [
                # chroma vector compact representation
                "chroma_top_idx", "chroma_top_value",
            ]
            ws2.append(snap_headers)
            for cell in ws2[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center

            snap_qs = (
                AudioFeatureSnapshot.objects
                .select_related("trackId")
                .order_by("trackId_id", "-syncedAt")
            )
            seen_tracks: set = set()
            for s in snap_qs:
                if s.trackId_id in seen_tracks:
                    continue  # keep only latest per track
                seen_tracks.add(s.trackId_id)
                snap = s.snapshot or {}
                fv = snap.get("featureVector") or {}
                chroma = fv.get("chroma") or []
                top_idx, top_val = "", ""
                if chroma:
                    top_idx = int(max(range(len(chroma)), key=lambda i: chroma[i]))
                    top_val = round(float(chroma[top_idx]), 4)
                ws2.append([
                    str(s.trackId_id),
                    (s.trackId.title if s.trackId else "") or "",
                    str(s.id),
                    _fmt_dt(s.syncedAt),
                    snap.get("calibrated", ""),
                    snap.get("energy"), snap.get("valence"),
                    snap.get("tempoBpm"), snap.get("acousticness"),
                    snap.get("instrumentalness"), snap.get("loudness"),
                    snap.get("keySignature"),
                    snap.get("primaryMood"), snap.get("moodConfidence"),
                    snap.get("genre"), snap.get("genreConfidence"),
                    snap.get("region"), snap.get("regionConfidence"),
                    snap.get("durationSec"),
                    *[fv.get(k) for k in _SNAPSHOT_DSP_KEYS],
                    top_idx, top_val,
                ])
            self._autofit(ws2, snap_headers, get_column_letter)
            self.stdout.write(
                f"  Snapshots: {len(seen_tracks)} tracks with at least one snapshot"
            )

        # -------------------- Summary sheet --------------------
        ws3 = wb.create_sheet("Summary")
        ws3.append(["section", "metric", "value"])
        for cell in ws3[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        ws3.append(["catalogue", "active_tracks", len(tracks)])
        ws3.append([
            "catalogue", "with_engine_output",
            sum(1 for t in tracks if t.energy is not None),
        ])
        ws3.append([
            "catalogue", "with_feature_snapshot",
            sum(1 for t in tracks if t.featuresSyncedAt is not None),
        ])
        ws3.append(["catalogue", "is_instrumental", sum(1 for t in tracks if t.isInstrumental)])

        # Per-numeric-feature stats
        for feature in _NUMERIC_FEATURES:
            values = [
                getattr(t, feature) for t in tracks
                if getattr(t, feature) is not None
            ]
            if not values:
                continue
            ws3.append(["feature", f"{feature}_n", len(values)])
            ws3.append(["feature", f"{feature}_mean", round(stats.fmean(values), 4)])
            ws3.append(["feature", f"{feature}_median", round(stats.median(values), 4)])
            ws3.append(["feature", f"{feature}_stdev", round(stats.pstdev(values), 4)])
            ws3.append(["feature", f"{feature}_min", round(min(values), 4)])
            ws3.append(["feature", f"{feature}_max", round(max(values), 4)])

        for label, key in (
            ("by_source", "source"),
            ("by_language", "language"),
            ("by_mood", "primaryMood"),
            ("by_genre", "genre"),
            ("by_region", "region"),
        ):
            counter: Counter = Counter()
            for t in tracks:
                counter[getattr(t, key) or "(unset)"] += 1
            for k, v in counter.most_common():
                ws3.append([label, str(k), v])

        # Out-of-range counters
        bad_counts: Counter = Counter()
        for t in tracks:
            for f, lo, hi in (
                ("energy", 0, 1), ("valence", 0, 1), ("acousticness", 0, 1),
                ("instrumentalness", 0, 1), ("loudness", -60, 0),
                ("tempoBpm", 40, 220),
            ):
                value = getattr(t, f)
                if value is None:
                    continue
                if not (lo <= value <= hi):
                    bad_counts[f] += 1
        for f, n in bad_counts.items():
            ws3.append(["out_of_range", f, n])
        if not bad_counts:
            ws3.append(["out_of_range", "(none)", 0])

        ws3.append(["meta", "generated_at",
                    datetime.now().isoformat(timespec="seconds")])
        self._autofit(ws3, ["section", "metric", "value"], get_column_letter)

        # -------------------- Save --------------------
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_dir = Path(settings.BASE_DIR) / "data" / "exports"
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = Path(options["path"]) if options["path"] else \
            out_dir / f"audio_audit_{ts}.xlsx"
        wb.save(out_path)

        self.stdout.write(self.style.SUCCESS(
            f"Wrote {out_path} — {len(tracks)} tracks, "
            f"{len(wb.sheetnames)} sheets"
        ))

    # ------------------------------------------------------------------
    @staticmethod
    def _autofit(ws, headers, get_column_letter) -> None:
        """Roughly auto-size column widths."""
        for col_idx in range(1, len(headers) + 1):
            letter = get_column_letter(col_idx)
            max_len = len(str(headers[col_idx - 1]))
            for row in ws.iter_rows(
                min_col=col_idx, max_col=col_idx, min_row=2,
                values_only=True,
            ):
                cell = row[0]
                if cell is None:
                    continue
                text = str(cell)
                if len(text) > max_len:
                    max_len = len(text)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 50)


def _fmt_dt(value) -> str:
    if value is None:
        return ""
    return value.isoformat(timespec="seconds") if hasattr(value, "isoformat") else str(value)
