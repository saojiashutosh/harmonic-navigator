"""Export all Artist and Track data to a shareable Excel file.

Usage (inside Docker):
    python manage.py export_to_excel
    python manage.py export_to_excel --output /app/data/harmonic_data.xlsx
"""
import os
from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from tracks.models import Artist, Track


ARTIST_COLS = [
    "id", "name", "spotify_id",
]

TRACK_COLS = [
    "id", "title", "artist_name", "artist_spotify_id",
    "type", "source", "spotify_id", "fma_id",
    "preview_url", "external_url", "stream_url",
    "release_year", "tempo_bpm", "duration_ms",
    "key_signature", "energy", "valence", "acousticness",
    "instrumentalness", "loudness",
    "primary_mood", "language", "genre", "region",
    "artist_popularity", "raga_name", "classical_form",
    "is_instrumental", "is_explicit", "is_active",
]

_HEADER_FILL = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
_HEADER_FONT = Font(color="E0C97F", bold=True)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=False)


def _style_header(ws):
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
    ws.row_dimensions[1].height = 20


class Command(BaseCommand):
    help = "Export all Artist and Track records to an Excel file (.xlsx)"

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            default="/app/harmonic_export.xlsx",
            help="Destination file path (default: /app/harmonic_export.xlsx)",
        )

    def handle(self, *args, **options):
        output_path = options["output"]
        wb = Workbook()

        # ── Sheet 1: Artists ───────────────────────────────────────────────
        ws_artists = wb.active
        ws_artists.title = "artists"
        ws_artists.append(ARTIST_COLS)
        _style_header(ws_artists)

        artists = Artist.objects.all().order_by("name")
        for a in artists.iterator():
            ws_artists.append([
                str(a.id),
                a.name or "",
                a.spotifyId or "",
            ])

        artist_count = artists.count()
        self.stdout.write(f"  Artists exported: {artist_count}")

        # ── Sheet 2: Tracks ────────────────────────────────────────────────
        ws_tracks = wb.create_sheet("tracks")
        ws_tracks.append(TRACK_COLS)
        _style_header(ws_tracks)

        tracks = (
            Track.objects.select_related("artistId")
            .all()
            .order_by("artistId__name", "title")
        )
        track_count = 0
        for t in tracks.iterator():
            ws_tracks.append([
                str(t.id),
                t.title or "",
                t.artistId.name if t.artistId else "",
                t.artistId.spotifyId if t.artistId else "",
                t.type or "",
                t.source or "",
                t.spotifyId or "",
                t.fmaId or "",
                t.previewUrl or "",
                t.externalUrl or "",
                t.streamUrl or "",
                t.releaseYear,
                t.tempoBpm,
                t.durationMs,
                t.keySignature or "",
                t.energy,
                t.valence,
                t.acousticness,
                t.instrumentalness,
                t.loudness,
                t.primaryMood or "",
                t.language or "",
                t.genre or "",
                t.region or "",
                t.artistPopularity,
                t.ragaName or "",
                t.classicalForm or "",
                t.isInstrumental,
                t.isExplicit,
                t.isActive,
            ])
            track_count += 1

        self.stdout.write(f"  Tracks exported:  {track_count}")

        # Auto-size columns (approximate — openpyxl has no built-in auto-fit)
        for ws in (ws_artists, ws_tracks):
            for col in ws.columns:
                max_len = max((len(str(cell.value or "")) for cell in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
        wb.save(output_path)
        self.stdout.write(self.style.SUCCESS(
            f"\nExported {artist_count} artists and {track_count} tracks → {output_path}"
        ))
