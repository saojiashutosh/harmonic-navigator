"""Import Artist and Track data from the exported Excel file into the DB.

This script is idempotent: it uses upsert logic so it is safe to run
multiple times without creating duplicate records.

Usage (inside Docker):
    python manage.py import_from_excel /app/harmonic_export.xlsx
    python manage.py import_from_excel /app/harmonic_export.xlsx --dry-run
"""
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from tracks.models import Artist, Track


def _str(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def _float(val):
    try:
        return float(val) if val is not None else None
    except (TypeError, ValueError):
        return None


def _int(val):
    try:
        return int(val) if val is not None else None
    except (TypeError, ValueError):
        return None


def _bool(val):
    if isinstance(val, bool):
        return val
    if isinstance(val, str):
        return val.strip().lower() in ("true", "1", "yes")
    return bool(val) if val is not None else False


class Command(BaseCommand):
    help = "Import Artist and Track records from the harmonic_export.xlsx file"

    def add_arguments(self, parser):
        parser.add_argument("excel_file", help="Path to the .xlsx file (inside the container)")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and validate the file without writing anything to the DB",
        )

    def handle(self, *args, **options):
        path = options["excel_file"]
        dry_run = options["dry_run"]

        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {path}")
        except Exception as exc:
            raise CommandError(f"Could not open file: {exc}")

        if "artists" not in wb.sheetnames:
            raise CommandError("Sheet 'artists' not found in the Excel file.")
        if "tracks" not in wb.sheetnames:
            raise CommandError("Sheet 'tracks' not found in the Excel file.")

        # ── Parse artists sheet ────────────────────────────────────────────
        ws_artists = wb["artists"]
        rows_artists = list(ws_artists.iter_rows(values_only=True))
        if not rows_artists:
            raise CommandError("artists sheet is empty.")
        artist_header = [str(h).strip() for h in rows_artists[0]]
        self.stdout.write(f"artists sheet — {len(rows_artists) - 1} rows")

        artists_to_upsert = []
        for row in rows_artists[1:]:
            if all(v is None for v in row):
                continue
            record = dict(zip(artist_header, row))
            artists_to_upsert.append({
                "id":         _str(record.get("id")),
                "name":       _str(record.get("name")),
                "spotify_id": _str(record.get("spotify_id")),
            })

        # ── Parse tracks sheet ─────────────────────────────────────────────
        ws_tracks = wb["tracks"]
        rows_tracks = list(ws_tracks.iter_rows(values_only=True))
        if not rows_tracks:
            raise CommandError("tracks sheet is empty.")
        track_header = [str(h).strip() for h in rows_tracks[0]]
        self.stdout.write(f"tracks sheet  — {len(rows_tracks) - 1} rows")

        tracks_to_upsert = []
        for row in rows_tracks[1:]:
            if all(v is None for v in row):
                continue
            record = dict(zip(track_header, row))
            tracks_to_upsert.append(record)

        wb.close()

        if dry_run:
            self.stdout.write(self.style.WARNING(
                f"\nDRY RUN — nothing written. "
                f"Parsed {len(artists_to_upsert)} artists and {len(tracks_to_upsert)} tracks."
            ))
            return

        # ── Write to DB ────────────────────────────────────────────────────
        with transaction.atomic():
            artist_created = artist_updated = 0

            # Build a name → Artist map for tracks that have no spotify_id
            artist_by_spotify: dict[str, Artist] = {}
            artist_by_name: dict[str, Artist] = {}

            for rec in artists_to_upsert:
                defaults = {"name": rec["name"]}
                if rec.get("spotify_id"):
                    artist, created = Artist.objects.update_or_create(
                        spotifyId=rec["spotify_id"],
                        defaults=defaults,
                    )
                else:
                    artist, created = Artist.objects.get_or_create(
                        name=rec["name"],
                        defaults=defaults,
                    )
                if created:
                    artist_created += 1
                else:
                    artist_updated += 1
                if rec.get("spotify_id"):
                    artist_by_spotify[rec["spotify_id"]] = artist
                if rec.get("name"):
                    artist_by_name[rec["name"]] = artist

            self.stdout.write(
                f"  Artists — created: {artist_created}, updated: {artist_updated}"
            )

            track_created = track_updated = track_skipped = 0

            for rec in tracks_to_upsert:
                artist_spotify_id = _str(rec.get("artist_spotify_id"))
                artist_name = _str(rec.get("artist_name"))

                artist = (
                    artist_by_spotify.get(artist_spotify_id)
                    if artist_spotify_id
                    else None
                ) or artist_by_name.get(artist_name)

                if artist is None:
                    # Artist not in the exported sheet — try DB lookup
                    if artist_spotify_id:
                        artist = Artist.objects.filter(spotifyId=artist_spotify_id).first()
                    if artist is None and artist_name:
                        artist = Artist.objects.filter(name=artist_name).first()
                    if artist is None and artist_name:
                        artist = Artist.objects.create(name=artist_name, spotifyId=artist_spotify_id)
                        artist_created += 1
                        artist_by_name[artist_name] = artist

                if artist is None:
                    self.stderr.write(
                        f"  Skipping track '{rec.get('title')}' — could not resolve artist '{artist_name}'"
                    )
                    track_skipped += 1
                    continue

                defaults = {
                    "title":            _str(rec.get("title")),
                    "artistId":         artist,
                    "type":             _str(rec.get("type")) or Track.TypeChoices.SONG,
                    "source":           _str(rec.get("source")) or Track.SourceChoices.MANUAL,
                    "fmaId":            _str(rec.get("fma_id")),
                    "previewUrl":       _str(rec.get("preview_url")),
                    "externalUrl":      _str(rec.get("external_url")),
                    "streamUrl":        _str(rec.get("stream_url")),
                    "releaseYear":      _int(rec.get("release_year")),
                    "tempoBpm":         _int(rec.get("tempo_bpm")),
                    "durationMs":       _int(rec.get("duration_ms")),
                    "keySignature":     _str(rec.get("key_signature")),
                    "energy":           _float(rec.get("energy")),
                    "valence":          _float(rec.get("valence")),
                    "acousticness":     _float(rec.get("acousticness")),
                    "instrumentalness": _float(rec.get("instrumentalness")),
                    "loudness":         _float(rec.get("loudness")),
                    "primaryMood":      _str(rec.get("primary_mood")),
                    "language":         _str(rec.get("language")),
                    "genre":            _str(rec.get("genre")),
                    "region":           _str(rec.get("region")),
                    "artistPopularity": _int(rec.get("artist_popularity")),
                    "ragaName":         _str(rec.get("raga_name")),
                    "classicalForm":    _str(rec.get("classical_form")),
                    "isInstrumental":   _bool(rec.get("is_instrumental")),
                    "isExplicit":       _bool(rec.get("is_explicit")),
                    "isActive":         _bool(rec.get("is_active", True)),
                }

                spotify_id = _str(rec.get("spotify_id"))
                if spotify_id:
                    _, created = Track.objects.update_or_create(
                        spotifyId=spotify_id,
                        defaults=defaults,
                    )
                else:
                    # No spotify_id — match on title + artist
                    title = defaults.get("title")
                    if title:
                        _, created = Track.objects.update_or_create(
                            title=title,
                            artistId=artist,
                            defaults=defaults,
                        )
                    else:
                        track_skipped += 1
                        continue

                if created:
                    track_created += 1
                else:
                    track_updated += 1

            self.stdout.write(
                f"  Tracks  — created: {track_created}, updated: {track_updated}, skipped: {track_skipped}"
            )

        self.stdout.write(self.style.SUCCESS("\nImport complete."))
