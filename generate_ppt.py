"""Generate Harmonic Navigator pitch deck (PPTX).

Color/font scheme mirrors the website's "Sand" paper palette + Editorial type.
Run: python generate_ppt.py
Outputs: HarmonicNavigator.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree


# ── Palette: Sand (matches .og-shell.palette-sand) ──────────────────────────
PAPER       = RGBColor(0xF0, 0xE6, 0xD2)
PAPER_2     = RGBColor(0xE8, 0xDF, 0xCC)
PAPER_3     = RGBColor(0xDC, 0xD0, 0xB8)
PAPER_FAB   = RGBColor(0xFF, 0xF8, 0xEC)
INK         = RGBColor(0x3A, 0x32, 0x2A)
INK_SOFT    = RGBColor(0x6F, 0x63, 0x54)
INK_FAINT   = RGBColor(0x9C, 0x90, 0x80)
SAGE        = RGBColor(0x7A, 0x86, 0x69)
ACCENT      = RGBColor(0xC2, 0x6F, 0x3C)
RULE        = RGBColor(0xCB, 0xC0, 0xAC)

FONT_DISPLAY = "Fraunces"
FONT_BODY = "Inter"
FONT_MONO = "JetBrains Mono"


# 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ── Helpers ─────────────────────────────────────────────────────────────────

def set_bg(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, *,
                font=FONT_BODY, size=18, bold=False, italic=False,
                color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
                line_spacing=1.2):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return tb


def add_rect(slide, left, top, width, height, *,
             fill=PAPER_2, line=None, line_width=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        if line_width is not None:
            shape.line.width = line_width
    shape.shadow.inherit = False
    return shape


def add_pill(slide, left, top, width, height, text, *,
             fill=PAPER_2, text_color=INK_SOFT, font=FONT_DISPLAY,
             size=12, italic=True):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = fill
    pill.line.color.rgb = RULE
    pill.line.width = Pt(0.75)
    pill.shadow.inherit = False
    tf = pill.text_frame
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.color.rgb = text_color
    return pill


def add_hline(slide, left, top, width, *, color=RULE, weight=0.75, dash=True):
    line = slide.shapes.add_connector(1, left, top, left + width, top)
    line.line.color.rgb = color
    line.line.width = Pt(weight)
    if dash:
        ln = line.line._get_or_add_ln()
        prstDash = etree.SubElement(ln, qn('a:prstDash'))
        prstDash.set('val', 'dash')
    return line


def add_dot(slide, cx, cy, r=Inches(0.06), color=ACCENT):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, r * 2, r * 2)
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    dot.shadow.inherit = False
    return dot


def add_eyebrow(slide, left, top, text, *, color=ACCENT):
    add_textbox(slide, left, top, Inches(6), Inches(0.4), text,
                font=FONT_DISPLAY, size=14, italic=True, color=color)


def add_logo(slide):
    add_dot(slide, Inches(0.55), Inches(0.55), r=Inches(0.07), color=SAGE)
    add_textbox(slide, Inches(0.7), Inches(0.32), Inches(4.5), Inches(0.5),
                "Harmonic Navigator.", font=FONT_DISPLAY, size=20,
                italic=True, color=INK)


def add_page_meta(slide, page_num, total, label):
    add_textbox(slide, Inches(9.8), Inches(0.4), Inches(3.3), Inches(0.4),
                f"§ {page_num:02d} / {total:02d}   ·   {label}",
                font=FONT_MONO, size=10, color=INK_FAINT, align=PP_ALIGN.RIGHT)


def add_footer(slide):
    add_hline(slide, Inches(0.55), Inches(7.05), Inches(12.2))
    add_textbox(slide, Inches(0.55), Inches(7.12), Inches(8), Inches(0.3),
                "harmonic navigator — a field guide to mood-led listening",
                font=FONT_DISPLAY, size=10, italic=True, color=INK_FAINT)
    add_textbox(slide, Inches(8.55), Inches(7.12), Inches(4.2), Inches(0.3),
                "2026",
                font=FONT_MONO, size=9, color=INK_FAINT, align=PP_ALIGN.RIGHT)


def _set_solidfill_alpha(shape, alpha_pct):
    """Apply transparency to a shape's solid fill. alpha_pct: 0 (invisible) – 100 (opaque)."""
    sp = shape.fill._xPr
    solidFill = sp.find(qn('a:solidFill'))
    if solidFill is None:
        return
    srgb = solidFill.find(qn('a:srgbClr'))
    if srgb is None:
        return
    # Remove any existing alpha first
    for old in srgb.findall(qn('a:alpha')):
        srgb.remove(old)
    alpha = etree.SubElement(srgb, qn('a:alpha'))
    alpha.set('val', str(int(alpha_pct * 1000)))  # OOXML uses 0–100000


def _set_line_alpha(shape, alpha_pct):
    """Apply transparency to a shape's line color."""
    ln = shape.line._get_or_add_ln()
    solidFill = ln.find(qn('a:solidFill'))
    if solidFill is None:
        return
    srgb = solidFill.find(qn('a:srgbClr'))
    if srgb is None:
        return
    for old in srgb.findall(qn('a:alpha')):
        srgb.remove(old)
    alpha = etree.SubElement(srgb, qn('a:alpha'))
    alpha.set('val', str(int(alpha_pct * 1000)))


def _send_to_back(shape):
    spTree = shape._element.getparent()
    spTree.remove(shape._element)
    spTree.insert(2, shape._element)


def add_ambient_blob(slide, cx, cy, r, *, color, alpha=5):
    """Soft, very-low-opacity ambient wash — mimics radial-gradient blobs."""
    blob = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, r * 2, r * 2)
    blob.fill.solid()
    blob.fill.fore_color.rgb = color
    blob.line.fill.background()
    blob.shadow.inherit = False
    _set_solidfill_alpha(blob, alpha)
    _send_to_back(blob)
    return blob


def add_bg_decor(slide):
    """Soft ambient washes — same vibe as the frontend's radial-gradient blobs."""
    add_ambient_blob(slide, Inches(1.5),  Inches(1.2), Inches(3.2),
                     color=SAGE, alpha=8)
    add_ambient_blob(slide, Inches(11.8), Inches(6.6), Inches(3.6),
                     color=ACCENT, alpha=8)
    add_ambient_blob(slide, Inches(12.6), Inches(2.4), Inches(2.4),
                     color=SAGE, alpha=6)
    add_ambient_blob(slide, Inches(0.6),  Inches(5.8), Inches(2.6),
                     color=ACCENT, alpha=6)


def chrome(slide, page_num, total, section_label):
    set_bg(slide)
    add_bg_decor(slide)
    add_logo(slide)
    add_page_meta(slide, page_num, total, section_label)
    add_footer(slide)


# ── Build deck ──────────────────────────────────────────────────────────────

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank = prs.slide_layouts[6]

TOTAL = 11


# ─────────────────────────── Slide 1 — The Question ─────────────────────────
s = prs.slides.add_slide(blank)
set_bg(s)
add_bg_decor(s)
add_logo(s)
add_page_meta(s, 1, TOTAL, "the question")
add_footer(s)

add_pill(s, Inches(0.55), Inches(1.7), Inches(2.3), Inches(0.4),
         "·  a thought experiment", fill=PAPER_2, text_color=SAGE)

add_textbox(s, Inches(0.55), Inches(2.25), Inches(12.2), Inches(1.2),
            "how often do you press play —",
            font=FONT_DISPLAY, size=54, color=INK)
add_textbox(s, Inches(0.55), Inches(3.25), Inches(12.2), Inches(1.2),
            "and skip, skip, skip?",
            font=FONT_DISPLAY, size=54, italic=True, color=ACCENT)

add_textbox(s, Inches(0.55), Inches(4.55), Inches(11.5), Inches(1.5),
            "The same song you loved last week doesn't fit today.\n"
            "Your library is huge — yet nothing matches the room you're in,\n"
            "the weather outside, or the mood inside your head.",
            font=FONT_DISPLAY, size=20, color=INK_SOFT, line_spacing=1.35)


# ─────────────────────────── Slide 2 — Why this matters ─────────────────────
s = prs.slides.add_slide(blank)
chrome(s, 2, TOTAL, "the problem")
add_eyebrow(s, Inches(0.55), Inches(1.5), "— so we asked —")

add_textbox(s, Inches(0.55), Inches(2.0), Inches(12.2), Inches(1.2),
            "what if the playlist",
            font=FONT_DISPLAY, size=44, color=INK)
add_textbox(s, Inches(0.55), Inches(2.85), Inches(12.2), Inches(1.2),
            "knew the mood first?",
            font=FONT_DISPLAY, size=44, italic=True, color=ACCENT)

add_hline(s, Inches(0.55), Inches(4.2), Inches(12.2))

cols = [
    ("01", "shuffle is random",
     "Algorithms guess from history.\nThey miss the state you're in\nright now."),
    ("02", "moods aren't genres",
     "Sad isn't a genre. Focused isn't\na genre. But they're how we\nactually listen."),
    ("03", "the moment is fleeting",
     "By the time you've curated the\nperfect mix, the mood has\nmoved on."),
]
col_w = Inches(3.95)
gap = Inches(0.15)
x0 = Inches(0.55)
for i, (num, title, body) in enumerate(cols):
    x = x0 + (col_w + gap) * i
    add_textbox(s, x, Inches(4.45), col_w, Inches(0.35),
                num, font=FONT_DISPLAY, size=14, italic=True, color=ACCENT)
    add_textbox(s, x, Inches(4.8), col_w, Inches(0.6),
                title, font=FONT_DISPLAY, size=22, color=INK)
    add_textbox(s, x, Inches(5.5), col_w, Inches(1.2),
                body, font=FONT_BODY, size=12, color=INK_SOFT, line_spacing=1.45)


# ─────────────────────────── Slide 3 — Introducing ──────────────────────────
s = prs.slides.add_slide(blank)
chrome(s, 3, TOTAL, "introducing")
add_eyebrow(s, Inches(0.55), Inches(1.5), "— introducing —")

add_textbox(s, Inches(0.55), Inches(2.0), Inches(12.2), Inches(1.6),
            "Harmonic",
            font=FONT_DISPLAY, size=88, italic=True, color=INK)
add_textbox(s, Inches(0.55), Inches(3.4), Inches(12.2), Inches(1.6),
            "Navigator.",
            font=FONT_DISPLAY, size=88, italic=True, color=ACCENT)

add_hline(s, Inches(0.55), Inches(5.4), Inches(12.2))

add_textbox(s, Inches(0.55), Inches(5.65), Inches(12.2), Inches(1.2),
            "Tell us how you feel — in ten small questions.\n"
            "We translate that into a mood, and curate a playlist tuned to it.",
            font=FONT_DISPLAY, size=20, color=INK_SOFT, line_spacing=1.4)


# ─────────────────────────── Slide 4 — Project overview / 10 Qs ─────────────
s = prs.slides.add_slide(blank)
chrome(s, 4, TOTAL, "project overview")
add_eyebrow(s, Inches(0.55), Inches(1.5), "— what we built —")
add_textbox(s, Inches(0.55), Inches(1.95), Inches(12.2), Inches(0.9),
            "ten questions →", font=FONT_DISPLAY, size=34, color=INK)
add_textbox(s, Inches(0.55), Inches(2.6), Inches(12.2), Inches(0.9),
            "one mood, one playlist.",
            font=FONT_DISPLAY, size=34, italic=True, color=ACCENT)

add_hline(s, Inches(0.55), Inches(3.55), Inches(12.2))

# The authentic 10 questions from QUESTION_DEFINITIONS
questions = [
    ("01", "energy",      "How's your energy feeling right now?"),
    ("02", "emotion",     "How are you feeling emotionally right now?"),
    ("03", "headspace",   "What is your headspace like right now?"),
    ("04", "activity",    "What are you currently doing?"),
    ("05", "company",     "Who are you with right now?"),
    ("06", "language",    "Which language songs do you prefer?"),
    ("07", "the goal",    "What should this playlist do for you?"),
    ("08", "an artist",   "Any artist you want to hear? (optional)"),
    ("09", "time of day", "What time of day is it for you?"),
    ("10", "era",         "From which era do you want your songs?"),
]
# 2-column list, 5 rows each
col_w_q = Inches(6.1)
row_h = Inches(0.55)
y0 = Inches(3.8)
for i, (num, tag, text) in enumerate(questions):
    r, c = divmod(i, 5)
    # 5 per column, fill left col first
    col = i // 5
    row = i % 5
    x = Inches(0.55) + col * (col_w_q + Inches(0.15))
    yy = y0 + row * row_h
    add_textbox(s, x, yy, Inches(0.45), row_h,
                num, font=FONT_MONO, size=11, color=ACCENT)
    add_textbox(s, x + Inches(0.5), yy, Inches(1.45), row_h,
                tag, font=FONT_DISPLAY, size=14, italic=True, color=SAGE)
    add_textbox(s, x + Inches(1.95), yy, col_w_q - Inches(2), row_h,
                text, font=FONT_BODY, size=12, color=INK)


# ─────────────────────────── Slide 5 — How it works ─────────────────────────
s = prs.slides.add_slide(blank)
chrome(s, 5, TOTAL, "how it works")
add_eyebrow(s, Inches(0.55), Inches(1.5), "— the flow —")

add_textbox(s, Inches(0.55), Inches(1.95), Inches(12.2), Inches(0.9),
            "from question to playlist —",
            font=FONT_DISPLAY, size=32, color=INK)
add_textbox(s, Inches(0.55), Inches(2.55), Inches(12.2), Inches(0.9),
            "in under a minute.",
            font=FONT_DISPLAY, size=32, italic=True, color=ACCENT)

steps = [
    ("01.", "ask",
     "Ten short prompts —\nyour energy, emotion,\nactivity, context."),
    ("02.", "infer",
     "Answers add up to\na mood label with\na confidence score."),
    ("03.", "curate",
     "Tracks pulled by\nmood, language, era —\nwith variety baked in."),
    ("04.", "play",
     "A bouquet you can\nplay, save, share —\nor re-pick anytime."),
]
card_w = Inches(2.95)
card_h = Inches(2.4)
card_gap = Inches(0.1)
y = Inches(3.85)
for i, (num, title, body) in enumerate(steps):
    x = Inches(0.55) + (card_w + card_gap) * i
    add_rect(s, x, y, card_w, card_h, fill=PAPER_2, line=RULE, line_width=Pt(0.75))
    add_textbox(s, x + Inches(0.25), y + Inches(0.2), card_w, Inches(0.4),
                num, font=FONT_DISPLAY, size=14, italic=True, color=ACCENT)
    add_textbox(s, x + Inches(0.25), y + Inches(0.55), card_w, Inches(0.6),
                title, font=FONT_DISPLAY, size=26, color=INK)
    add_textbox(s, x + Inches(0.25), y + Inches(1.25), card_w - Inches(0.5), Inches(1.1),
                body, font=FONT_BODY, size=11, color=INK_SOFT, line_spacing=1.5)
    if i < len(steps) - 1:
        ax = x + card_w + Inches(0.01)
        add_textbox(s, ax, y + Inches(1.0), card_gap + Inches(0.1), Inches(0.4),
                    "→", font=FONT_DISPLAY, size=18, color=ACCENT,
                    align=PP_ALIGN.CENTER)


# ─────────────────────────── Slide 6 — The mood engine (simplified) ─────────
s = prs.slides.add_slide(blank)
chrome(s, 6, TOTAL, "the engine")
add_eyebrow(s, Inches(0.55), Inches(1.5), "— the heart of it —")
add_textbox(s, Inches(0.55), Inches(1.95), Inches(12.2), Inches(0.9),
            "the mood engine.",
            font=FONT_DISPLAY, size=38, italic=True, color=INK)
add_textbox(s, Inches(0.55), Inches(2.7), Inches(12.2), Inches(0.6),
            "Every answer nudges eight possible moods. The strongest one wins.",
            font=FONT_DISPLAY, size=18, italic=True, color=INK_SOFT)

add_hline(s, Inches(0.55), Inches(3.55), Inches(12.2))

# 8 mood labels — the real ones
add_textbox(s, Inches(0.55), Inches(3.75), Inches(12.2), Inches(0.4),
            "the eight moods we listen for",
            font=FONT_DISPLAY, size=14, italic=True, color=INK_SOFT)

moods = ["celebratory", "energized", "focused", "calm",
         "melancholic", "anxious", "reflective", "tender"]
gx, gy = Inches(0.55), Inches(4.2)
mw, mh = Inches(2.95), Inches(0.55)
for i, m in enumerate(moods):
    r, c = divmod(i, 4)
    x = gx + (mw + Inches(0.1)) * c
    yy = gy + (mh + Inches(0.15)) * r
    add_rect(s, x, yy, mw, mh, fill=PAPER_FAB, line=RULE, line_width=Pt(0.5))
    add_textbox(s, x + Inches(0.2), yy + Inches(0.1), mw, mh,
                m, font=FONT_DISPLAY, size=16, italic=True, color=INK)
    add_textbox(s, x, yy + Inches(0.1), mw - Inches(0.25), mh,
                "·", font=FONT_DISPLAY, size=15, color=ACCENT, align=PP_ALIGN.RIGHT)

add_textbox(s, Inches(0.55), Inches(5.95), Inches(12.2), Inches(0.5),
            "Emotion and energy get the loudest vote. Time-of-day and company whisper.",
            font=FONT_DISPLAY, size=14, italic=True, color=INK_SOFT)


# ─────────────────────────── Slide 7 — How we weight the answers ────────────
s = prs.slides.add_slide(blank)
chrome(s, 7, TOTAL, "the weighting")
add_eyebrow(s, Inches(0.55), Inches(1.5), "— under the hood —")
add_textbox(s, Inches(0.55), Inches(1.95), Inches(12.2), Inches(0.9),
            "not all questions",
            font=FONT_DISPLAY, size=34, color=INK)
add_textbox(s, Inches(0.55), Inches(2.55), Inches(12.2), Inches(0.9),
            "weigh the same.",
            font=FONT_DISPLAY, size=34, italic=True, color=ACCENT)

add_hline(s, Inches(0.55), Inches(3.5), Inches(12.2))

# Left: category multipliers table
add_textbox(s, Inches(0.55), Inches(3.7), Inches(6.0), Inches(0.4),
            "category multipliers",
            font=FONT_DISPLAY, size=14, italic=True, color=INK_SOFT)

cat_rows = [
    ("emotion",     "1.60×", "emotional tone"),
    ("energy",      "1.30×", "energy level"),
    ("cognition",   "1.10×", "headspace"),
    ("activity",    "1.00×", "what you're doing"),
    ("preference",  "0.90×", "language · goal · era · artist"),
    ("context",     "0.70×", "company · time of day"),
]
tx = Inches(0.55)
ty = Inches(4.15)
tw = Inches(6.0)
rh = Inches(0.38)
add_rect(s, tx, ty, tw, rh * len(cat_rows) + Inches(0.1),
         fill=PAPER_FAB, line=RULE, line_width=Pt(0.5))
for i, (cat, mult, qs) in enumerate(cat_rows):
    yy = ty + Inches(0.05) + rh * i
    add_textbox(s, tx + Inches(0.2), yy + Inches(0.04), Inches(1.4), rh,
                cat, font=FONT_DISPLAY, size=14, italic=True, color=INK)
    add_textbox(s, tx + Inches(1.6), yy + Inches(0.04), Inches(0.9), rh,
                mult, font=FONT_MONO, size=13, color=ACCENT)
    add_textbox(s, tx + Inches(2.5), yy + Inches(0.04), tw - Inches(2.7), rh,
                qs, font=FONT_BODY, size=11, color=INK_SOFT)
    if i < len(cat_rows) - 1:
        add_hline(s, tx + Inches(0.15), yy + rh,
                  tw - Inches(0.3), color=RULE, dash=True)

# Right: how the score is built — small explainer cards
right_x = Inches(7.0)
right_w = Inches(5.75)
ny = Inches(3.7)
add_textbox(s, right_x, ny, right_w, Inches(0.4),
            "how a mood score is built",
            font=FONT_DISPLAY, size=14, italic=True, color=INK_SOFT)

notes = [
    ("01.", "per-answer nudge",
     "Every option carries a hand-tuned\nweight for each of the moods —\nplaylist goal & emotion push hardest."),
    ("02.", "category multiplier",
     "Multiply that nudge by the table\non the left — emotion counts 2.3×\nas much as time-of-day."),
    ("03.", "synergy bonuses",
     "Certain combinations (e.g. sad +\ndrifting + escape) earn an extra\nbump toward a matching mood."),
    ("04.", "softmax & blend",
     "Sums become probabilities (×4\nsharpening). If top two are close,\nwe blend them into the playlist."),
]
ncols = 2
cw = (right_w - Inches(0.15)) / 2
ch = Inches(1.18)
n_y0 = ny + Inches(0.45)
for i, (num, title, body) in enumerate(notes):
    r, c = divmod(i, ncols)
    x = right_x + (cw + Inches(0.15)) * c
    yy = n_y0 + (ch + Inches(0.12)) * r
    add_rect(s, x, yy, cw, ch, fill=PAPER_2, line=RULE, line_width=Pt(0.5))
    add_textbox(s, x + Inches(0.18), yy + Inches(0.1), Inches(0.5), Inches(0.3),
                num, font=FONT_MONO, size=10, color=ACCENT)
    add_textbox(s, x + Inches(0.55), yy + Inches(0.08), cw - Inches(0.7), Inches(0.35),
                title, font=FONT_DISPLAY, size=14, italic=True, color=INK)
    add_textbox(s, x + Inches(0.18), yy + Inches(0.45), cw - Inches(0.35), Inches(0.7),
                body, font=FONT_BODY, size=10, color=INK_SOFT, line_spacing=1.4)

add_textbox(s, Inches(0.55), Inches(6.7), Inches(12.2), Inches(0.3),
            "Tell us how you feel — the engine listens to emotion loudest, then energy, then everything else.",
            font=FONT_DISPLAY, size=13, italic=True, color=INK_SOFT)


# ─────────────────────────── Slide 8 — Features (single slide) ──────────────
s = prs.slides.add_slide(blank)
chrome(s, 8, TOTAL, "features")
add_eyebrow(s, Inches(0.55), Inches(1.5), "— what's inside —")
add_textbox(s, Inches(0.55), Inches(1.95), Inches(12.2), Inches(0.9),
            "features.", font=FONT_DISPLAY, size=38, italic=True, color=INK)

feats = [
    ("mood survey",      "Ten quick prompts. Skippable\nartist field for your favourites."),
    ("smart inference",  "Maps your answers to one of\neight moods, with a confidence."),
    ("curated playlist", "Filtered by mood, language,\nand era — variety baked in."),
    ("re-pick anytime",  "Mood shifted? Get a fresh\nbouquet without re-doing the survey."),
    ("save & manage",    "My Playlists — name, edit, replay.\nAdd individual tracks on the fly."),
    ("group sessions",   "Friends scan a QR code. Moods\nmerge. The room gets one mix."),
]
gx, gy = Inches(0.55), Inches(3.05)
fw, fh = Inches(4.05), Inches(1.85)
for i, (title, body) in enumerate(feats):
    r, c = divmod(i, 3)
    x = gx + (fw + Inches(0.1)) * c
    yy = gy + (fh + Inches(0.15)) * r
    add_rect(s, x, yy, fw, fh, fill=PAPER_2, line=RULE, line_width=Pt(0.5))
    add_textbox(s, x + Inches(0.25), yy + Inches(0.12), Inches(1), Inches(0.3),
                f"{i+1:02d}", font=FONT_MONO, size=10, color=ACCENT)
    add_textbox(s, x + Inches(0.25), yy + Inches(0.42), fw - Inches(0.5), Inches(0.5),
                title, font=FONT_DISPLAY, size=20, color=INK)
    add_textbox(s, x + Inches(0.25), yy + Inches(1.0), fw - Inches(0.5), Inches(0.85),
                body, font=FONT_BODY, size=11, color=INK_SOFT, line_spacing=1.45)


# ─────────────────────────── Slide 9 — Tech stack (light) ───────────────────
s = prs.slides.add_slide(blank)
chrome(s, 9, TOTAL, "tech stack")
add_eyebrow(s, Inches(0.55), Inches(1.5), "— built with —")
add_textbox(s, Inches(0.55), Inches(1.95), Inches(12.2), Inches(0.9),
            "tech stack.", font=FONT_DISPLAY, size=38, italic=True, color=INK)

pillars = [
    ("frontend",
     ["React 18", "Vite", "qrcode.react"],
     "Field-guide UI — paper palettes,\nhand-set type, responsive layout."),
    ("backend",
     ["Django 5.2", "Django REST Framework", "PostgreSQL", "Redis · django-redis", "Gunicorn"],
     "REST API, persistence, and\ncached mood inference."),
    ("data & auth",
     ["Spotipy (Spotify API)", "scikit-learn · NumPy · pandas", "librosa", "Firebase Admin", "cryptography"],
     "Track catalogue, mood scoring,\nsigned auth, encrypted secrets."),
]
px = Inches(0.55)
pw = Inches(4.05)
gap = Inches(0.1)
py = Inches(2.95)
ph = Inches(3.55)
for i, (title, techs, body) in enumerate(pillars):
    x = px + (pw + gap) * i
    add_rect(s, x, py, pw, ph, fill=PAPER_2, line=RULE, line_width=Pt(0.6))
    add_textbox(s, x + Inches(0.3), py + Inches(0.22), pw, Inches(0.5),
                title, font=FONT_DISPLAY, size=22, italic=True, color=ACCENT)
    add_hline(s, x + Inches(0.3), py + Inches(0.85), pw - Inches(0.6),
              color=RULE, dash=True)
    # Tech list — bullet rows
    iy = py + Inches(1.0)
    for t in techs:
        add_dot(s, x + Inches(0.42), iy + Inches(0.13), r=Inches(0.035), color=SAGE)
        add_textbox(s, x + Inches(0.6), iy, pw - Inches(0.8), Inches(0.32),
                    t, font=FONT_MONO, size=11, color=INK)
        iy += Inches(0.32)
    # Description below the list
    add_textbox(s, x + Inches(0.3), py + ph - Inches(0.95),
                pw - Inches(0.6), Inches(0.85),
                body, font=FONT_BODY, size=11, italic=True,
                color=INK_SOFT, line_spacing=1.45)

add_textbox(s, Inches(0.55), Inches(6.7), Inches(12.2), Inches(0.3),
            "All wrapped in Docker Compose — one command spins up the whole app.",
            font=FONT_DISPLAY, size=14, italic=True, color=INK_SOFT)


# ─────────────────────────── Slide 10 — Demo time ───────────────────────────
s = prs.slides.add_slide(blank)
set_bg(s)
add_bg_decor(s)
add_logo(s)
add_page_meta(s, 10, TOTAL, "demo")
add_footer(s)

# Centered big "demo time"
add_pill(s, Inches(5.55), Inches(2.0), Inches(2.3), Inches(0.45),
         "·  live walkthrough", fill=PAPER_2, text_color=SAGE)

add_textbox(s, Inches(0.55), Inches(2.85), Inches(12.2), Inches(1.6),
            "demo time.",
            font=FONT_DISPLAY, size=96, italic=True, color=ACCENT,
            align=PP_ALIGN.CENTER)

add_hline(s, Inches(2.55), Inches(4.95), Inches(8.2))

add_textbox(s, Inches(0.55), Inches(5.2), Inches(12.2), Inches(0.6),
            "ten questions · one mood · a fresh playlist.",
            font=FONT_DISPLAY, size=20, italic=True, color=INK_SOFT,
            align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.55), Inches(5.85), Inches(12.2), Inches(0.4),
            "follow along — pick a mood, see what plays.",
            font=FONT_DISPLAY, size=14, italic=True, color=INK_FAINT,
            align=PP_ALIGN.CENTER)


# ─────────────────────────── Slide 11 — Questions / Suggestions ─────────────
s = prs.slides.add_slide(blank)
set_bg(s)
add_bg_decor(s)
add_logo(s)
add_page_meta(s, 11, TOTAL, "over to you")

add_textbox(s, Inches(0.55), Inches(2.4), Inches(12.2), Inches(1.4),
            "questions",
            font=FONT_DISPLAY, size=72, color=INK)
add_textbox(s, Inches(0.55), Inches(3.5), Inches(12.2), Inches(1.4),
            "& suggestions.",
            font=FONT_DISPLAY, size=72, italic=True, color=ACCENT)

add_hline(s, Inches(0.55), Inches(5.4), Inches(12.2))

add_textbox(s, Inches(0.55), Inches(5.6), Inches(8), Inches(0.5),
            "Harmonic Navigator",
            font=FONT_DISPLAY, size=18, italic=True, color=INK_SOFT)

add_textbox(s, Inches(8.55), Inches(5.6), Inches(4.2), Inches(0.5),
            "thank you.",
            font=FONT_DISPLAY, size=22, italic=True, color=INK,
            align=PP_ALIGN.RIGHT)
add_textbox(s, Inches(8.55), Inches(6.05), Inches(4.2), Inches(0.4),
            "we'd love to hear from you.",
            font=FONT_DISPLAY, size=13, italic=True, color=INK_SOFT,
            align=PP_ALIGN.RIGHT)

add_footer(s)


# ── Save ────────────────────────────────────────────────────────────────────
import os
out = "HarmonicNavigator.pptx"
try:
    prs.save(out)
except PermissionError:
    out = "HarmonicNavigator_v2.pptx"
    prs.save(out)
print(f"saved: {out}  ·  {len(prs.slides)} slides")
